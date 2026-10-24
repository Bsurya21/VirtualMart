
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

PROJECT_DIR = Path(__file__).resolve().parent.parent
OUTPUT_FILE = PROJECT_DIR / "docs" / "VirtualMart_Report.docx"
ERD_FILE = PROJECT_DIR / "diagrams" / "VirtualMart_ERD.png"

doc = Document()

# Page setup and document styles
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10)
styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(28)
styles["Title"].font.color.rgb = RGBColor(31, 78, 121)

for style_name in ("Heading 1", "Heading 2"):
    styles[style_name].font.name = "Aptos Display"
    styles[style_name].font.color.rgb = RGBColor(31, 78, 121)


def add_bullet(text):
    doc.add_paragraph(text, style="List Bullet")


def add_number(text):
    doc.add_paragraph(text, style="List Number")


def add_table(headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Shading Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = str(header)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True

    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    doc.add_paragraph()
    return table


# Cover page
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("\n\n")

title = doc.add_heading("VirtualMart", 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle = doc.add_paragraph("Database Modeling and Architecture Design")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.runs[0].font.size = Pt(17)

p = doc.add_paragraph(
    "\nWeek 1 Project Report\n"
    "Technology: PostgreSQL 18\n"
    "Environment: Local development database\n"
    "Dataset: Fictional synthetic demonstration data"
)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_page_break()

# 1. Executive summary
doc.add_heading("1. Executive Summary", level=1)
doc.add_paragraph(
    "VirtualMart is a sample e-commerce database project designed to demonstrate "
    "relational data modeling, database architecture, integrity constraints, indexing, "
    "and analytical SQL. The database uses PostgreSQL and organizes customers, "
    "products, suppliers, inventory, orders, payments, shipments, returns, and "
    "inventory activity into related tables."
)
doc.add_paragraph(
    "All sample records are fictional and intended only for demonstration. "
    "The project is configured for local development and does not require personal "
    "datasets, external accounts, or private API credentials."
)

# 2. Objectives
doc.add_heading("2. Project Objectives", level=1)
for item in [
    "Design a relational schema for a fictional online marketplace.",
    "Represent business entities and their relationships using primary and foreign keys.",
    "Apply data types, CHECK constraints, and uniqueness rules to improve data quality.",
    "Create indexes for common lookup and join operations.",
    "Load synthetic records and test analytical SQL queries.",
    "Document the schema using an entity-relationship diagram (ERD).",
]:
    add_bullet(item)

# 3. Technology stack
doc.add_heading("3. Technology Stack and Environment", level=1)
add_table(
    ["Component", "Purpose"],
    [
        ["PostgreSQL 18", "Relational database engine"],
        ["pgAdmin 4", "Database administration and SQL query execution"],
        ["Visual Studio Code", "SQL scripts and project documentation"],
        ["diagrams.net", "ERD diagram creation"],
        ["Python and python-docx", "Local generation of the Word report"],
    ],
)

doc.add_paragraph(
    "The database is named virtualmart_db, and the project schema is virtualmart. "
    "SQL scripts are maintained in the sql directory. The project also contains "
    "diagrams, docs, and scripts directories."
)

# 4. Architecture
doc.add_heading("4. Database Architecture", level=1)
doc.add_paragraph(
    "The database follows a relational design. Entity tables hold core business "
    "records, while relationship tables represent many-to-many associations or "
    "transaction details. Foreign keys connect related records and help prevent "
    "orphaned references."
)
add_bullet("Schema: virtualmart")
add_bullet("Database: virtualmart_db")
add_bullet("Database engine: PostgreSQL 18")
add_bullet("Sample records are synthetic and stored locally.")

doc.add_heading("4.1 Entity-Relationship Diagram", level=2)
if ERD_FILE.exists():
    doc.add_paragraph(
        "The following diagram provides a visual overview of the 16-table model."
    )
    doc.add_picture(str(ERD_FILE), width=Inches(6.4))
else:
    doc.add_paragraph(
        "The ERD is maintained separately in diagrams/VirtualMart_ERD.drawio "
        "and diagrams/VirtualMart_ERD.png. The image was not embedded because "
        "the expected PNG file was not found when this report was generated."
    )

doc.add_paragraph(
    "The diagram is an overview of the model. The SQL schema definitions remain "
    "the authoritative reference for exact column definitions and constraints."
)

# 5. Tables
doc.add_heading("5. Database Tables", level=1)
tables = [
    ("customers", "Customer contact and profile records."),
    ("customer_addresses", "Customer delivery or billing addresses."),
    ("categories", "Product category definitions."),
    ("suppliers", "Supplier records."),
    ("products", "Product catalogue, pricing, and product attributes."),
    ("supplier_products", "Associations between suppliers and products."),
    ("warehouses", "Warehouse or stock-location records."),
    ("inventory", "Product stock by warehouse, including reserved quantities."),
    ("sales_orders", "Order headers, status, addresses, and monetary totals."),
    ("order_items", "Products and quantities contained in each order."),
    ("payments", "Payment references, methods, statuses, and amounts."),
    ("shipments", "Shipment and tracking information for orders."),
    ("returns", "Return request or return header records."),
    ("return_items", "Products and quantities associated with returns."),
    ("inventory_movements", "Recorded inventory changes and movement history."),
    ("audit_log", "A foundation for application-recorded audit events."),
]
add_table(["Table", "Purpose"], tables)

# 6. Integrity and modeling
doc.add_heading("6. Data Modeling and Integrity", level=1)
for item in [
    "Primary keys uniquely identify records in each table.",
    "Foreign keys represent relationships between customers, addresses, products, orders, and related records.",
    "Identity columns provide generated numeric identifiers where configured.",
    "CHECK constraints validate selected values and business rules.",
    "UNIQUE constraints protect selected identifiers from duplication.",
    "NUMERIC columns are used for monetary amounts to avoid floating-point rounding behavior.",
    "Timestamp columns record relevant dates and times.",
    "The inventory table uses a composite key for its product and warehouse relationship.",
]:
    add_bullet(item)

doc.add_paragraph(
    "The schema is structured to reduce unnecessary duplication and separate "
    "entities from transaction details. Formal normalization and constraint "
    "coverage should be assessed against the SQL definitions before production use."
)

# 7. Indexes
doc.add_heading("7. Indexing Strategy", level=1)
doc.add_paragraph(
    "The project includes 11 explicitly created indexes for commonly used "
    "lookup and join paths. PostgreSQL also creates indexes automatically for "
    "primary-key and unique constraints where required. The reported index count "
    "in pgAdmin can therefore be higher than the number of indexes explicitly "
    "defined in the indexing script."
)
doc.add_paragraph(
    "Indexes can improve query performance, but they also consume storage and "
    "add overhead to writes. Production index choices should be validated with "
    "representative workloads and EXPLAIN ANALYZE."
)

# 8. Sample data
doc.add_heading("8. Synthetic Sample Dataset", level=1)
doc.add_paragraph(
    "The initial demonstration dataset contains 53 rows across the 16 tables. "
    "This includes 5 customers, 3 categories, 3 suppliers, 4 products, "
    "2 warehouses, 3 sales orders, 3 order items, 3 payments, 2 shipments, "
    "8 inventory records, and 8 inventory movements. Returns and audit_log "
    "currently contain no sample rows."
)
doc.add_paragraph(
    "The sample data is intentionally small to make query results easy to inspect. "
    "It is not intended to represent production-scale volume."
)

# 9. Analytics
doc.add_heading("9. Analytical SQL and Test Results", level=1)
doc.add_paragraph(
    "Eight analytical queries were executed and reviewed in pgAdmin. They cover:"
)
for item in [
    "Order counts and monetary totals grouped by order status.",
    "Top-selling products by quantity and line-item revenue.",
    "Customer order history, including customers with no orders.",
    "Inventory reorder alerts based on available stock.",
    "Payment monitoring by status and payment method.",
    "Order shipment and tracking status.",
    "Supplier coverage across products.",
    "Daily order volume and value.",
]:
    add_bullet(item)

doc.add_heading("9.1 Verified Demonstration Results", level=2)
add_table(
    ["Order", "Status", "Total (INR)"],
    [
        ["VM-ORD-1001", "DELIVERED", "1727.90"],
        ["VM-ORD-1002", "SHIPPED", "1198.95"],
        ["VM-ORD-1003", "PROCESSING", "668.95"],
        ["Total", "3 orders", "3595.80"],
    ],
)
doc.add_paragraph(
    "The inventory reorder query returned no rows because all 8 inventory records "
    "had 100 units on hand, 0 reserved, and a reorder level of 10. This is an "
    "expected result for the current synthetic dataset, not evidence that the "
    "reorder-alert query is unnecessary."
)

# 10. Security
doc.add_heading("10. Security and Privacy Considerations", level=1)
for item in [
    "Only fictional demonstration records are used.",
    "No personal API keys, passwords, or production credentials are included in project files.",
    "Database access should use least-privilege roles when the project is extended.",
    "Credentials should not be committed to source control or stored in SQL scripts.",
    "The audit_log table is not currently wired to automatic triggers; automated auditing would require additional implementation and testing.",
    "Backups, access controls, retention rules, and monitoring would need to be configured before production use.",
]:
    add_bullet(item)

# 11. Limitations
doc.add_heading("11. Limitations and Assumptions", level=1)
for item in [
    "The dataset is small, synthetic, and intended for learning and demonstration.",
    "The ERD is a visual overview; SQL scripts define the authoritative schema.",
    "The audit_log table is a foundation only and does not automatically capture database changes.",
    "Returns and return_items have been created but are not populated in the current sample dataset.",
    "Performance has not been benchmarked with production-sized data.",
    "The analytics results reflect only the current sample records.",
]:
    add_bullet(item)

# 12. Future enhancements
doc.add_heading("12. Future Enhancements", level=1)
for item in [
    "Add more synthetic orders, products, inventory changes, and return scenarios.",
    "Add automated integrity checks for order totals, payment amounts, and inventory quantities.",
    "Implement and test controlled audit logging if required.",
    "Review query plans with EXPLAIN ANALYZE and benchmark indexes on larger datasets.",
    "Add repeatable test scripts and a documented database reset procedure.",
    "Refine the ERD with explicit cardinalities and complete key relationships.",
]:
    add_bullet(item)

# 13. Project files
doc.add_heading("13. Project Structure", level=1)
for item in [
    "sql/01_schema.sql — schema setup",
    "sql/02_tables.sql — table definitions",
    "sql/03_indexes.sql — explicit indexes",
    "sql/04_sample_data.sql — synthetic demonstration data",
    "sql/05_analytics.sql — analytical SELECT queries",
    "diagrams/VirtualMart_ERD.drawio — editable ERD",
    "diagrams/VirtualMart_ERD.png — exported ERD image",
    "docs/VirtualMart_Report.docx — this report",
    "scripts/generate_report.py — local report-generation script",
]:
    add_bullet(item)

doc.add_heading("14. Conclusion", level=1)
doc.add_paragraph(
    "The VirtualMart project demonstrates the foundational workflow of designing "
    "a relational database, defining entities and relationships, adding integrity "
    "constraints and indexes, inserting synthetic data, and validating analytical "
    "queries. The project is a useful starting point for further work in database "
    "development and data analysis. Additional automated testing, security controls, "
    "and performance evaluation would be required before any production deployment."
)

# Footer on each section page
for sec in doc.sections:
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("VirtualMart | Database Modeling and Architecture Design")

OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT_FILE)
print(f"Report created successfully: {OUTPUT_FILE}")
